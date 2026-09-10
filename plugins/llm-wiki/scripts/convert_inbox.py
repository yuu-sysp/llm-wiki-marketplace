#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""inbox の非 .md ファイルを .md へ変換する.

ingest / lint が見るのは `inbox/*.md` だけなので、pptx・xlsx・docx・pdf などを
inbox に置いても黙って無視されてしまう。このスクリプトが変換して取り込み対象に載せる。

- 変換後の原本は `sources/_attachments/` へ移動して保持する（表や図を後から確認できる）。
- 生成 md に frontmatter は付けない（ingest Phase A が付与するので二重になる）。
- 変換ライブラリが無い場合はクラッシュせずスキップし、最後に pip コマンドを提示する。

vault ルートは argv[1] → 環境変数の順で解決。
  例: python convert_inbox.py "D:/資料/LLM-Wiki"
終了コード: 0 = 変換完了（対象なしも 0）／1 = スキップあり／2 = vault 未解決
"""
import csv as csv_mod
import io
import re
import shutil
import sys
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _vault import print_unresolved_hint, resolve_vault  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

TODAY = datetime.now().strftime("%Y-%m-%d")

# 打ち切り上限。超えた場合は黙って切らず、md と標準出力の両方に明記する。
MAX_XLSX_ROWS = 2000
MAX_PDF_PAGES = 300

TEXT_EXT = {".txt", ".log", ".json", ".yml", ".yaml"}
TABLE_EXT = {".csv", ".tsv"}
HTML_EXT = {".html", ".htm"}
LEGACY_EXT = {".doc", ".xls", ".ppt"}          # 旧バイナリ形式（非対応）

PIP_HINT = "pip install python-docx openpyxl python-pptx pypdf"


class MissingLib(Exception):
    """変換ライブラリが未導入。パッケージ名を持つ。"""

    def __init__(self, package):
        super().__init__(package)
        self.package = package


# ---------------------------------------------------------------- 共通ヘルパー

def read_text_guess(path: Path) -> str:
    """日本語テキストは cp932 のことが多いので順に試す。"""
    for enc in ("utf-8-sig", "utf-8", "cp932"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def cell(value) -> str:
    """markdown table のセルに入れられる形に均す。"""
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>").strip()


def md_table(rows) -> str:
    """2次元リスト → markdown table（先頭行をヘッダ扱い）。"""
    rows = [r for r in rows]
    while rows and all(cell(c) == "" for c in rows[-1]):     # 末尾の空行を落とす
        rows.pop()
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [list(r) + [None] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(cell(c) for c in rows[0]) + " |",
           "|" + "|".join([" --- "] * width) + "|"]
    for r in rows[1:]:
        out.append("| " + " | ".join(cell(c) for c in r) + " |")
    return "\n".join(out)


def unique_path(directory: Path, stem: str, suffix: str) -> Path:
    """既存ファイルを絶対に壊さない。衝突したら -1, -2 … と採番する。"""
    cand = directory / f"{stem}{suffix}"
    n = 1
    while cand.exists():
        cand = directory / f"{stem}-{n}{suffix}"
        n += 1
    return cand


# ---------------------------------------------------------------- 各形式の変換
# 各コンバータは (本文, 概要ラベル) を返す。ライブラリが無ければ MissingLib を投げる。

def conv_text(path: Path):
    body = read_text_guess(path).replace("\r\n", "\n")
    return body, f"{len(body.splitlines())} 行"


def conv_table(path: Path):
    text = read_text_guess(path)
    delim = "\t" if path.suffix.lower() == ".tsv" else ","
    rows = list(csv_mod.reader(io.StringIO(text), delimiter=delim))
    return md_table(rows), f"{len(rows)} 行"


class _Text(HTMLParser):
    """script/style を除いた本文テキストだけ拾う。"""

    def __init__(self):
        super().__init__()
        self.parts, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip += 1
        elif tag in ("p", "br", "div", "tr", "li", "h1", "h2", "h3", "h4"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self.skip:
            self.skip -= 1

    def handle_data(self, data):
        if not self.skip and data.strip():
            self.parts.append(data.strip())


def conv_html(path: Path):
    p = _Text()
    p.feed(read_text_guess(path))
    body = re.sub(r"\n{3,}", "\n\n", " ".join(p.parts).replace(" \n ", "\n")).strip()
    return body, f"{len(body.splitlines())} 行"


def conv_docx(path: Path):
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        raise MissingLib("python-docx")

    doc = docx.Document(str(path))
    out, n_tbl = [], 0
    # 段落と表の順序を保つため body の子要素を順に辿る
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            m = re.search(r"(\d+)", style)
            if ("Heading" in style or "見出し" in style) and m:
                level = min(int(m.group(1)), 5) + 1      # ページ見出しと衝突させない
                out.append("#" * level + " " + text)
            elif style.startswith("Title") or style.startswith("表題"):
                out.append("## " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            rows = [[c.text for c in row.cells] for row in Table(child, doc).rows]
            table = md_table(rows)
            if table:
                out.append(table)
                n_tbl += 1

    label = f"{len([o for o in out if o])} ブロック"
    if n_tbl:
        label += f" / 表 {n_tbl}"
    return "\n\n".join(out), label


def conv_xlsx(path: Path):
    try:
        import openpyxl
    except ImportError:
        raise MissingLib("openpyxl")

    # data_only=True で数式ではなく計算済みの値を取る
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out, total, truncated = [], 0, []
    n_sheets = 0
    try:
        n_sheets = len(wb.worksheets)              # close 後は参照できないのでここで取る
        for ws in wb.worksheets:
            rows, over = [], False
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_XLSX_ROWS:
                    over = True
                    break
                rows.append(list(row))
            total += len(rows)
            out.append(f"## {ws.title}")
            table = md_table(rows)
            out.append(table if table else "_（空のシート）_")
            if over:
                truncated.append(ws.title)
                out.append(f"> ⚠ このシートは {MAX_XLSX_ROWS} 行で打ち切りました。"
                           f"全量は原本 `{path.name}` を参照してください。")
    finally:
        wb.close()

    label = f"{n_sheets} シート / {total} 行"
    if truncated:
        label += f" / 打ち切り: {', '.join(truncated)}"
    return "\n\n".join(out), label


def conv_pptx(path: Path):
    try:
        from pptx import Presentation
    except ImportError:
        raise MissingLib("python-pptx")

    prs = Presentation(str(path))
    out, n, n_tbl = [], 0, 0
    for i, slide in enumerate(prs.slides, 1):
        n = i
        out.append(f"## Slide {i}")
        texts, tables = [], []
        for shape in slide.shapes:
            # 表は GraphicFrame で has_text_frame が False。拾わないと
            # 「表や図を後から確認できる」という取り込みの前提が崩れる。
            if getattr(shape, "has_table", False):
                rows = [[c.text for c in row.cells] for row in shape.table.rows]
                table = md_table(rows)
                if table:
                    tables.append(table)
                    n_tbl += 1
                continue
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                t = "".join(run.text for run in para.runs).strip()
                if t:
                    texts.append(t)
        if texts:
            out.append("\n".join(texts))
        out.extend(tables)
        if not texts and not tables:
            out.append("_（テキストなし）_")

        if slide.has_notes_slide:
            note = (slide.notes_slide.notes_text_frame.text or "").strip()
            if note:
                out.append("> **ノート**: " + note.replace("\n", "\n> "))
    label = f"{n} スライド"
    if n_tbl:
        label += f" / 表 {n_tbl}"
    return "\n\n".join(out), label


def conv_pdf(path: Path):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise MissingLib("pypdf")

    reader = PdfReader(str(path))
    out, over = [], False
    for i, page in enumerate(reader.pages, 1):
        if i > MAX_PDF_PAGES:
            over = True
            break
        text = (page.extract_text() or "").strip()
        out.append(f"## p.{i}\n\n" + (text if text else "_（テキストを抽出できませんでした）_"))

    label = f"{min(len(reader.pages), MAX_PDF_PAGES)} ページ"
    if over:
        out.append(f"> ⚠ {MAX_PDF_PAGES} ページで打ち切りました（全 {len(reader.pages)} ページ）。"
                   f"残りは原本 `{path.name}` を参照してください。")
        label += f" / 打ち切り（全 {len(reader.pages)}）"
    return "\n\n".join(out), label


def converter_for(suffix: str):
    if suffix in TEXT_EXT:
        return conv_text
    if suffix in TABLE_EXT:
        return conv_table
    if suffix in HTML_EXT:
        return conv_html
    return {".docx": conv_docx,
            ".xlsx": conv_xlsx, ".xlsm": conv_xlsx,
            ".pptx": conv_pptx,
            ".pdf": conv_pdf}.get(suffix)


# ---------------------------------------------------------------------- main

def main() -> int:
    vault = resolve_vault()
    if vault is None:
        print_unresolved_hint("convert_inbox.py")
        return 2
    if not vault.is_dir():
        print(f"ERROR: vault が存在しません: {vault}")
        return 2

    inbox = vault / "inbox"
    if not inbox.is_dir():
        print(f"inbox がありません: {inbox}（bootstrap.py を先に実行してください）")
        return 2

    entries = [p for p in inbox.iterdir() if not p.name.startswith(".")]
    targets = sorted(p for p in entries
                     if p.is_file()
                     and p.suffix.lower() != ".md")
    # フォルダは再帰しない（どの階層を1ページ相当と見なすかが決められないため）。
    # 黙って無視すると「置いたのに取り込まれない」になるので、スキップとして必ず出す。
    dirs = sorted(p.name for p in entries if p.is_dir())
    if not targets and not dirs:
        return 0                                   # 対象なしは静かに終了（毎回走るため）

    attachments = vault / "sources" / "_attachments"
    attachments.mkdir(parents=True, exist_ok=True)

    print(f"LLM-Wiki convert-inbox  (vault: {vault})")
    converted, skipped, missing = [], [], set()
    for d in dirs:
        skipped.append((d + "/", "フォルダは対象外。中のファイルを inbox 直下へ出してください"))

    for src in targets:
        suffix = src.suffix.lower()
        fn = converter_for(suffix)

        if fn is None:
            reason = ("旧バイナリ形式は非対応。.docx / .xlsx / .pptx で保存し直してください"
                      if suffix in LEGACY_EXT else "未対応の拡張子")
            skipped.append((src.name, reason))
            continue

        try:
            body, label = fn(src)
        except MissingLib as e:
            missing.add(e.package)
            skipped.append((src.name, f"{e.package} が未導入"))
            continue
        except Exception as e:                     # 壊れたファイル等で全体を止めない
            skipped.append((src.name, f"変換失敗: {type(e).__name__}: {e}"))
            continue

        # 原本を先に退避し、md からは確定したファイル名で参照する
        dest = unique_path(attachments, src.stem, src.suffix)
        shutil.move(str(src), str(dest))

        header = (f"> **変換元**: `{src.name}`（{suffix.lstrip('.')} / {label}）\n"
                  f"> **原本**: `sources/_attachments/{dest.name}`\n"
                  f"> **変換日**: {TODAY} — `convert_inbox.py` が自動生成\n")
        out_md = unique_path(inbox, src.stem, ".md")
        out_md.write_text(f"# {src.stem}\n\n{header}\n{body}\n", encoding="utf-8")

        converted.append((src.name, out_md.name, label))

    for name, md, label in converted:
        print(f"  + {name}  ->  {md}  ({label})")
    for name, reason in skipped:
        print(f"  - {name}  スキップ: {reason}")

    print(f"\n変換: {len(converted)} 件  /  スキップ: {len(skipped)} 件")
    if converted:
        print(f"原本は sources/_attachments/ に保存しました（{len(converted)} 件）")
    if missing:
        print("\n変換ライブラリが不足しています。次を実行してください:")
        print(f"  {PIP_HINT}")
        print(f"  （不足: {', '.join(sorted(missing))}）")
        print("  ※ PDF はライブラリが無くても、ingest 時に Claude が直接読み取れます。")

    return 1 if skipped else 0


if __name__ == "__main__":
    sys.exit(main())
